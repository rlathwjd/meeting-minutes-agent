import shutil
from tempfile import TemporaryDirectory
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any, Dict, List, Optional

from docx import Document

from .word_converter import WordConversionError, convert_doc_to_docx, convert_docx_to_doc


def format_meeting_datetime(value: str) -> str:
    meeting_datetime = datetime.fromisoformat(value.replace("Z", "+00:00"))
    return meeting_datetime.strftime("%Y-%m-%d %H:%M")


def format_meeting_datetime_for_template(value: str) -> str:
    weekdays = ["월", "화", "수", "목", "금", "토", "일"]
    meeting_datetime = datetime.fromisoformat(value.replace("Z", "+00:00"))
    meeting_datetime = meeting_datetime.astimezone(timezone(timedelta(hours=9)))
    weekday = weekdays[meeting_datetime.weekday()]
    return meeting_datetime.strftime(f"%Y. %m. %d({weekday}) %H:%M")


def format_meeting_type(value: str) -> str:
    return "대면" if value == "in_person" else "비대면"


def format_meeting_type_for_template(value: str) -> str:
    return "대면 회의" if value == "in_person" else "비대면 회의"


def format_attendees(payload: Dict[str, Any], *, markdown: bool) -> str:
    attendees = payload.get("attendees_by_company", [])
    lines = []
    for row in attendees:
        company = row.get("company", "").strip()
        names = ", ".join(row.get("attendees", []))
        if company or names:
            prefix = "- " if markdown else ""
            lines.append(f"{prefix}{company}: {names}")

    if lines:
        return "\n".join(lines)
    return "- 참석자 정보 없음" if markdown else "참석자 정보 없음"


def format_list(items: List[str], *, empty: str) -> str:
    lines = [f"- {item}" for item in items if item]
    return "\n".join(lines) if lines else empty


def format_action_items(items: List[Dict[str, str]]) -> str:
    lines = []
    for item in items:
        task = item.get("task", "").strip()
        owner = item.get("owner", "").strip()
        due_date = item.get("due_date", "").strip()
        details = []
        if owner:
            details.append(f"담당: {owner}")
        if due_date:
            details.append(f"기한: {due_date}")
        suffix = f" ({', '.join(details)})" if details else ""
        if task:
            lines.append(f"- {task}{suffix}")

    return "\n".join(lines) if lines else "후속조치 없음"


def build_minutes_text(
    payload: Dict[str, Any],
    transcript_text: str,
    minutes_content: Optional[Dict[str, Any]] = None,
) -> str:
    content = minutes_content or {}
    title = content.get("title") or payload.get("manual_title") or "회의록"
    summary = content.get("summary") or transcript_text.strip() or "녹취 내용이 비어 있습니다."
    agenda = format_list(content.get("agenda", []), empty="안건 없음")
    decisions = format_list(content.get("decisions", []), empty="결정사항 없음")
    action_items = format_action_items(content.get("action_items", []))
    notes = format_list(content.get("notes", []), empty="특이사항 없음")

    return "\n".join(
        [
            f"# {title}",
            "",
            f"- 회의 일시: {format_meeting_datetime(payload['meeting_datetime'])}",
            f"- 회의 장소: {payload.get('location', '')}",
            f"- 작성자: {payload.get('author', '')}",
            f"- 회의 형식: {format_meeting_type(payload.get('meeting_type', 'in_person'))}",
            "",
            "## 참석자",
            format_attendees(payload, markdown=True),
            "",
            "## 요약",
            summary,
            "",
            "## 안건",
            agenda,
            "",
            "## 결정사항",
            decisions,
            "",
            "## 후속조치",
            action_items,
            "",
            "## 특이사항",
            notes,
            "",
        ]
    )


def build_placeholder_values(
    payload: Dict[str, Any],
    transcript_text: str,
    minutes_content: Optional[Dict[str, Any]] = None,
) -> Dict[str, str]:
    content = minutes_content or {}
    summary = content.get("summary") or transcript_text.strip()
    agenda = format_list(content.get("agenda", []), empty="")
    decisions = format_list(content.get("decisions", []), empty="")
    action_items = format_action_items(content.get("action_items", []))
    notes = format_list(content.get("notes", []), empty="")

    return {
        "{{title}}": content.get("title") or payload.get("manual_title") or "회의록",
        "{{meeting_datetime}}": format_meeting_datetime(payload["meeting_datetime"]),
        "{{location}}": payload.get("location", ""),
        "{{author}}": payload.get("author", ""),
        "{{meeting_type}}": format_meeting_type(payload.get("meeting_type", "in_person")),
        "{{attendees}}": format_attendees(payload, markdown=False),
        "{{summary}}": summary,
        "{{agenda}}": agenda,
        "{{decisions}}": decisions,
        "{{action_items}}": action_items,
        "{{notes}}": notes,
        "{{meeting_content}}": summary,
        "{{transcript}}": transcript_text.strip(),
        "{{minutes}}": build_minutes_text(payload, transcript_text, minutes_content),
    }


def build_meeting_content_for_template(
    transcript_text: str,
    minutes_content: Optional[Dict[str, Any]] = None,
) -> str:
    content = minutes_content or {}
    sections = []
    if content.get("summary"):
        sections.append(f"1. 주요 내용\n{content['summary']}")
    if content.get("agenda"):
        sections.append("2. 안건\n" + "\n".join(f"- {item}" for item in content["agenda"] if item))
    if content.get("decisions"):
        sections.append("3. 결정사항\n" + "\n".join(f"- {item}" for item in content["decisions"] if item))
    if content.get("notes"):
        sections.append("4. 특이사항\n" + "\n".join(f"- {item}" for item in content["notes"] if item))

    return "\n\n".join(sections) if sections else transcript_text.strip() or "녹취 내용이 비어 있습니다."


def set_cell_text(table, row_index: int, cell_index: int, value: str) -> None:
    if row_index >= len(table.rows) or cell_index >= len(table.rows[row_index].cells):
        return
    table.rows[row_index].cells[cell_index].text = value


def fill_known_minutes_template(
    document: Document,
    payload: Dict[str, Any],
    transcript_text: str,
    minutes_content: Optional[Dict[str, Any]] = None,
) -> bool:
    if len(document.tables) < 3:
        return False

    table0, table1, table2 = document.tables[:3]
    if not table1.rows or "제" not in table1.rows[0].cells[0].text or "목" not in table1.rows[0].cells[0].text:
        return False
    if not table2.rows or "회" not in table2.rows[0].cells[-1].text or "내" not in table2.rows[0].cells[-1].text:
        return False

    content = minutes_content or {}
    title = content.get("title") or payload.get("manual_title") or "회의록"
    attendees = format_attendees(payload, markdown=False)
    meeting_content = build_meeting_content_for_template(transcript_text, minutes_content)
    action_items = format_action_items(content.get("action_items", []))

    set_cell_text(table0, 1, 2, payload.get("author", ""))
    set_cell_text(table1, 0, 1, title)
    set_cell_text(table1, 1, 1, format_meeting_datetime_for_template(payload["meeting_datetime"]))
    set_cell_text(table1, 1, 3, payload.get("location", ""))
    set_cell_text(table1, 2, 1, attendees)
    set_cell_text(table1, 3, 1, format_meeting_type_for_template(payload.get("meeting_type", "in_person")))
    set_cell_text(table2, 1, 1, meeting_content)
    set_cell_text(table2, 2, 1, action_items)
    return True


def replace_in_paragraph(paragraph, values: Dict[str, str]) -> bool:
    full_text = paragraph.text
    next_text = full_text
    for key, value in values.items():
        next_text = next_text.replace(key, value)

    if next_text == full_text:
        return False

    for run in paragraph.runs:
        run.text = ""
    if paragraph.runs:
        paragraph.runs[0].text = next_text
    else:
        paragraph.add_run(next_text)
    return True


def fill_docx_template(
    template_path: Path,
    output_path: Path,
    payload: Dict[str, Any],
    transcript_text: str,
    minutes_content: Optional[Dict[str, Any]] = None,
) -> None:
    document = Document(template_path)
    values = build_placeholder_values(payload, transcript_text, minutes_content)
    changed = False

    for paragraph in document.paragraphs:
        changed = replace_in_paragraph(paragraph, values) or changed

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    changed = replace_in_paragraph(paragraph, values) or changed

    if not changed:
        changed = fill_known_minutes_template(document, payload, transcript_text, minutes_content)

    if not changed:
        document.add_page_break()
        for line in build_minutes_text(payload, transcript_text, minutes_content).splitlines():
            document.add_paragraph(line)

    document.save(output_path)


def generate_minutes_file(
    template_path: Path,
    output_path: Path,
    payload: Dict[str, Any],
    transcript_text: str,
    minutes_content: Optional[Dict[str, Any]] = None,
) -> None:
    suffix = template_path.suffix.lower()
    if suffix == ".doc":
        with TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            converted_template = temp_path / "template.docx"
            converted_output = temp_path / "generated.docx"
            convert_doc_to_docx(template_path, converted_template)
            fill_docx_template(converted_template, converted_output, payload, transcript_text, minutes_content)
            convert_docx_to_doc(converted_output, output_path)
        return

    if suffix == ".docx":
        fill_docx_template(template_path, output_path, payload, transcript_text, minutes_content)
        return

    if suffix in {".txt", ".md"}:
        output_path.write_text(build_minutes_text(payload, transcript_text, minutes_content), encoding="utf-8")
        return

    shutil.copyfile(template_path, output_path)
